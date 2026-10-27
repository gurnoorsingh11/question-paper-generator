from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

def export_paper(metadata, sections, paper, instructions=""):
    doc = Document()
    heading = doc.add_heading(metadata.get("college_name") or
                              "Swami Vivekanand Institute of Engineering & Technology", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(metadata.get("programme_semester", "B. Tech."))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(metadata.get("subject_name", "Subject"))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(
        f"Subject Code: {metadata.get('subject_code','')}    "
        f"M. Code: {metadata.get('m_code','')}"
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        f"Time: {metadata.get('duration','180')} minutes    "
        f"Max. Marks: {metadata.get('maximum_marks','')}"
    )
    if metadata.get("exam_date"):
        doc.add_paragraph(f"Date of Examination: {metadata['exam_date']}")
    doc.add_heading("INSTRUCTIONS TO CANDIDATES", level=1)
    doc.add_paragraph(instructions or "Read all instructions carefully before attempting the paper.")
    q_number = 1
    for section in paper:
        doc.add_heading(section["section"].upper(), level=1)
        for q in section["questions"]:
            p = doc.add_paragraph()
            p.add_run(f"{q_number}. {q['question_text']} ").bold = True
            p.add_run(f"[{section['marks_each']} marks]")
            q_number += 1
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream.getvalue()
