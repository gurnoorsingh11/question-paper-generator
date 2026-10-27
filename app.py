import io
import random
import re
from datetime import date

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

st.set_page_config(page_title="SVIET Question Paper Generator", page_icon="📄", layout="wide")

COLLEGE_DEFAULT = "Swami Vivekanand Institute of Engineering & Technology"
UMC_LINE = (
    "NOTE: Disclosure of Identity by writing Mobile No. or Marking of passing request "
    "on any paper of Answer Sheet will lead to UMC against the Student."
)
MARK_OPTIONS = [2, 4, 5, 8, 10]

st.title("SVIET Question Paper Generator")
st.caption("Build a question paper in the same overall layout as the supplied Sample QP.docx.")


def normalize(value):
    return re.sub(r"\s+", " ", str(value)).strip()


def looks_like_heading(text):
    value = normalize(text).strip(" :-")
    if not value:
        return True
    return bool(re.match(
        r"^(unit\s+[ivx\d]+|module\s+\d+|chapter\s+\d+|question\s*(bank|text)?|questions?|s\.?\s*no\.?|serial\s*no\.?)$",
        value, flags=re.IGNORECASE
    ))


def unique(items):
    result, seen = [], set()
    for item in items:
        text = normalize(item)
        if len(text) < 3 or looks_like_heading(text):
            continue
        key = text.casefold()
        if key not in seen:
            seen.add(key)
            result.append(text)
    return result


def read_docx(upload):
    upload.seek(0)
    doc = Document(upload)
    items = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                items.append(max(cells, key=len))
    return unique(items)


def read_excel(upload):
    upload.seek(0)
    workbook = pd.ExcelFile(upload)
    items = []
    for sheet_name in workbook.sheet_names:
        frame = pd.read_excel(workbook, sheet_name=sheet_name, header=None).fillna("")
        for _, row in frame.iterrows():
            cells = [normalize(v) for v in row.tolist() if normalize(v)]
            if cells:
                items.append(max(cells, key=len))
    return unique(items)


def read_bank(upload):
    name = upload.name.lower()
    if name.endswith(".docx"):
        return read_docx(upload)
    if name.endswith((".xlsx", ".xls")):
        return read_excel(upload)
    raise ValueError("Upload a Word (.docx) or Excel (.xlsx/.xls) file.")


def add_paragraph(doc, text="", *, bold=False, size=10, align=None, space_after=0, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(str(text))
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    return p


def number_to_words(n):
    words = {
        0: "ZERO", 1: "ONE", 2: "TWO", 3: "THREE", 4: "FOUR", 5: "FIVE",
        6: "SIX", 7: "SEVEN", 8: "EIGHT", 9: "NINE", 10: "TEN",
        11: "ELEVEN", 12: "TWELVE", 13: "THIRTEEN", 14: "FOURTEEN",
        15: "FIFTEEN", 16: "SIXTEEN", 17: "SEVENTEEN", 18: "EIGHTEEN",
        19: "NINETEEN", 20: "TWENTY"
    }
    return words.get(int(n), str(int(n)))


def add_centered_heading(doc, text):
    add_paragraph(doc, text, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)


def build_docx(meta, section_a, section_b, section_c, marks_b, marks_c, attempt_count):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)

    # Sample QP header and title block
    add_paragraph(doc, "Roll No. __________________________________", size=10, space_after=0)
    add_paragraph(doc, f"Total No. of Questions: {1 + len(section_b) + len(section_c):02d}", size=10)
    add_paragraph(doc, meta["college"], bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_paragraph(doc, meta["degree_semester"], bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_paragraph(doc, meta["subject"], bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_paragraph(doc, f"Subject Code: {meta['subject_code']}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    if meta["m_code"].strip():
        add_paragraph(doc, f"M. Code: {meta['m_code']}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, f"Date of Examination: {meta['exam_date']}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(
        doc,
        f"Time: {meta['duration']}\t\t\t\t\t\tMax. Marks: {meta['maximum_marks']}",
        size=10,
    )

    add_paragraph(doc, "INSTRUCTIONS TO CANDIDATES:", bold=True, size=10, space_after=2)
    add_paragraph(
        doc,
        f"SECTION-A is COMPULSORY consisting of {number_to_words(len(section_a))} questions carrying TWO marks each.",
        size=10,
    )
    add_paragraph(
        doc,
        f"SECTION-B has {number_to_words(len(section_b))} questions and SECTION-C has {number_to_words(len(section_c))} questions.",
        size=10,
    )
    if marks_b == marks_c:
        marks_text = number_to_words(marks_b)
        add_paragraph(
            doc,
            f"Attempt any {number_to_words(min(attempt_count, len(section_b) + len(section_c)))} questions from SECTION B & C carrying {marks_text} marks each.",
            size=10,
        )
    else:
        add_paragraph(
            doc,
            f"SECTION-B questions carry {marks_b} marks each and SECTION-C questions carry {marks_c} marks each. Attempt questions as directed.",
            size=10,
        )
    add_paragraph(doc, "Select at least TWO questions from SECTION-B & C.", size=10, space_after=5)

    # Sections and question spaces follow the sample's centered headings.
    add_centered_heading(doc, "SECTION-A")
    add_paragraph(doc, "1. Write short notes on:", size=10, space_after=3)
    sub_labels = "abcdefghijklmnopqrstuvwxyz"
    for index, question in enumerate(section_a):
        label = sub_labels[index] if index < len(sub_labels) else str(index + 1)
        add_paragraph(doc, f"({label}) {question}", size=10, space_after=5)

    add_centered_heading(doc, "SECTION-B")
    q_number = 2
    for question in section_b:
        add_paragraph(doc, f"{q_number}. {question}  [{marks_b} marks]", size=10, space_after=8)
        q_number += 1

    add_centered_heading(doc, "SECTION-C")
    # Continue numbering after Section B, matching the sample paper's overall numbering.
    for question in section_c:
        add_paragraph(doc, f"{q_number}. {question}  [{marks_c} marks]", size=10, space_after=8)
        q_number += 1

    # UMC warning is deliberately the final line after all generated questions.
    add_paragraph(doc, UMC_LINE, bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


with st.expander("1. Exam details", expanded=True):
    left, right = st.columns(2)
    with left:
        college = st.text_input("College / institute", COLLEGE_DEFAULT)
        degree_semester = st.text_input("Course / semester heading", "B. Tech. (CSE) (Sem.–1)")
        subject = st.text_input("Subject name", "Engineering Physics")
        subject_code = st.text_input("Subject code", "BPH-101-26")
    with right:
        m_code = st.text_input("M. Code", "93794")
        exam_date = st.date_input("Date of examination", value=date.today())
        duration = st.text_input("Time / duration", "3 Hrs.")
        maximum_marks = st.number_input("Maximum marks", min_value=1, max_value=300, value=60, step=1)

st.subheader("2. Upload question banks")
col_a, col_long = st.columns(2)
with col_a:
    short_upload = st.file_uploader("Section A — 2-mark questions", type=["docx", "xlsx", "xls"], key="short_bank")
with col_long:
    long_upload = st.file_uploader("Sections B and C — long questions", type=["docx", "xlsx", "xls"], key="long_bank")


def load_if_changed(upload, state_key, display_name):
    if upload is None:
        st.session_state.pop(state_key, None)
        st.session_state.pop(state_key + "_fingerprint", None)
        return []
    fingerprint = f"{upload.name}:{upload.size}"
    if st.session_state.get(state_key + "_fingerprint") != fingerprint:
        try:
            st.session_state[state_key] = read_bank(upload)
            st.session_state[state_key + "_fingerprint"] = fingerprint
            st.session_state.pop("selected_paper", None)
        except Exception as exc:
            st.error(f"Couldn't read {display_name}: {exc}")
            return []
    return st.session_state.get(state_key, [])


short_questions = load_if_changed(short_upload, "short_questions", "the Section A question bank")
long_questions = load_if_changed(long_upload, "long_questions", "the Sections B/C question bank")
if short_upload:
    st.caption(f"Section A bank: {len(short_questions)} questions loaded.")
if long_upload:
    st.caption(f"Long-question bank: {len(long_questions)} questions loaded.")

st.subheader("3. Question counts and marks")
ca, cb, cc = st.columns(3)
with ca:
    count_a = st.number_input("Section A questions", min_value=0, max_value=200, value=10, step=1)
with cb:
    count_b = st.number_input("Section B questions", min_value=0, max_value=200, value=4, step=1)
with cc:
    count_c = st.number_input("Section C questions", min_value=0, max_value=200, value=4, step=1)
mb, mc = st.columns(2)
with mb:
    marks_b = st.selectbox("Marks per Section B question", MARK_OPTIONS, index=3)
with mc:
    marks_c = st.selectbox("Marks per Section C question", MARK_OPTIONS, index=3)
attempt_count = st.number_input("Questions to attempt from Sections B/C (for instructions)", min_value=0, max_value=200, value=5, step=1)

short_count = min(int(count_a), len(short_questions))
long_count_b = min(int(count_b), len(long_questions))
long_count_c = min(int(count_c), max(0, len(long_questions) - long_count_b))
if int(count_a) > len(short_questions):
    st.info(f"Section A requested {int(count_a)} questions; only {len(short_questions)} are available. All available questions will be used.")
if int(count_b) + int(count_c) > len(long_questions):
    st.info(f"Sections B and C request {int(count_b) + int(count_c)} long questions, but only {len(long_questions)} are available. Questions won't be duplicated across B and C.")


def choose_paper():
    a = random.sample(short_questions, min(int(count_a), len(short_questions)))
    pool = long_questions.copy()
    random.shuffle(pool)
    b = pool[:min(int(count_b), len(pool))]
    remaining = pool[len(b):]
    c = remaining[:min(int(count_c), len(remaining))]
    st.session_state["selected_paper"] = {"A": a, "B": b, "C": c}


act1, act2 = st.columns(2)
with act1:
    if st.button("🎲 Generate random paper", type="primary", use_container_width=True):
        if not short_questions and not long_questions:
            st.error("Upload at least one question bank first.")
        else:
            choose_paper()
            st.rerun()
with act2:
    if st.button("🔀 Shuffle all sections", use_container_width=True):
        if not short_questions and not long_questions:
            st.error("Upload at least one question bank first.")
        else:
            choose_paper()
            st.rerun()

selected = st.session_state.get("selected_paper")
if selected is not None:
    st.subheader("4. Preview and shuffle")
    st.markdown(f"**{college}**  \n{degree_semester}  \n**{subject}** | Code: {subject_code or '—'} | M. Code: {m_code or '—'}  \nDate: {exam_date.strftime('%d-%m-%Y')} | Time: {duration} | Max. Marks: {maximum_marks}")
    for key, title, requested in [("A", "SECTION-A — Write short notes on:", int(count_a)), ("B", f"SECTION-B ({marks_b} marks each)", int(count_b)), ("C", f"SECTION-C ({marks_c} marks each)", int(count_c))]:
        head_col, shuffle_col = st.columns([5, 1])
        with head_col:
            st.markdown(f"### {title}")
        with shuffle_col:
            if st.button("Shuffle", key=f"shuffle_{key}", use_container_width=True):
                if key == "A":
                    selected["A"] = random.sample(short_questions, min(int(count_a), len(short_questions)))
                else:
                    other = "C" if key == "B" else "B"
                    reserved = set(selected.get(other, []))
                    pool = [q for q in long_questions if q not in reserved]
                    selected[key] = random.sample(pool, min(requested, len(pool)))
                st.session_state["selected_paper"] = selected
                st.rerun()
        questions = selected.get(key, [])
        if not questions:
            st.caption("No questions selected for this section.")
        else:
            for i, question in enumerate(questions, start=1):
                mark_suffix = "" if key == "A" else f" [{marks_b if key == 'B' else marks_c} marks]"
                st.write(f"{i}. {question}{mark_suffix}")

    meta = {
        "college": college,
        "degree_semester": degree_semester,
        "subject": subject,
        "subject_code": subject_code,
        "m_code": m_code,
        "exam_date": exam_date.strftime("%d-%m-%Y"),
        "duration": duration,
        "maximum_marks": int(maximum_marks),
    }
    try:
        data = build_docx(meta, selected.get("A", []), selected.get("B", []), selected.get("C", []), int(marks_b), int(marks_c), int(attempt_count))
        safe_subject = re.sub(r"[^A-Za-z0-9_-]+", "_", subject).strip("_") or "SVIET_Exam"
        st.download_button(
            "⬇️ Download Word question paper",
            data=data,
            file_name=f"{safe_subject}_Question_Paper.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )
    except Exception as exc:
        st.error(f"Could not create the Word paper: {exc}")
